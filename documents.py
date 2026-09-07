"""Read source locations before splitting text, and carry them through retrieval."""

from dataclasses import dataclass, replace
from io import BytesIO
from pathlib import Path
import re


@dataclass(frozen=True)
class Passage:
    text: str
    pages: tuple[int, ...] = ()
    locations: tuple[str, ...] = ()


@dataclass(frozen=True)
class SourceChunk:
    text: str
    filename: str
    pages: tuple[int, ...] = ()
    locations: tuple[str, ...] = ()

    @property
    def label(self):
        if self.pages:
            location = "PDF 第 " + "、".join(map(str, self.pages)) + " 页"
        else:
            location = "；".join(self.locations)
        return f"{self.filename} · {location}"


@dataclass(frozen=True)
class Document:
    filename: str
    passages: tuple[Passage, ...]
    warnings: tuple[str, ...] = ()

    @property
    def text(self):
        return "\n\n".join(p.text for p in self.passages)


def _paragraphs(text):
    return [p.strip() for p in re.split(r"\n\s*\n", text.replace("\r\n", "\n").replace("\r", "\n")) if p.strip()]


def text_document(text, filename="输入文本"):
    return Document(filename, tuple(Passage(p, locations=(f"第 {i} 段",))
                                    for i, p in enumerate(_paragraphs(text), 1)))


def read_document(data: bytes, filename: str) -> Document:
    suffix = Path(filename).suffix.lower()
    passages, warnings = [], []
    try:
        if suffix == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(BytesIO(data))
            if reader.is_encrypted and not reader.decrypt(""):
                raise ValueError("PDF 受密码保护，请先解锁后上传。")
            for page_number, page in enumerate(reader.pages, 1):
                text = page.extract_text() or ""
                if not text.strip():
                    warnings.append(f"PDF 第 {page_number} 页未提取到文字，可能为空白页或扫描图片；该页未参与检索。")
                passages.extend(Passage(p, pages=(page_number,)) for p in _paragraphs(text))
        elif suffix == ".docx":
            import docx
            from docx.text.paragraph import Paragraph
            paragraph_number = table_number = 0
            for block in docx.Document(BytesIO(data)).iter_inner_content():
                if isinstance(block, Paragraph):
                    paragraph_number += 1
                    if block.text.strip():
                        passages.append(Passage(block.text.strip(), locations=(f"正文第 {paragraph_number} 段",)))
                else:
                    table_number += 1
                    for row_number, row in enumerate(block.rows, 1):
                        text = " | ".join(cell.text.strip() for cell in row.cells)
                        if any(cell.text.strip() for cell in row.cells):
                            passages.append(Passage(text, locations=(f"表格 {table_number} 第 {row_number} 行",)))
        elif suffix == ".txt":
            return text_document(data.decode("utf-8-sig"), filename)
        else:
            raise ValueError("仅支持 PDF、DOCX 和 UTF-8 TXT 文件。")
    except UnicodeDecodeError as exc:
        raise ValueError("TXT 编码无法识别，请另存为 UTF-8 后上传。") from exc
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("文档读取失败，请检查文件是否损坏或受密码保护。") from exc
    return Document(filename, tuple(passages), tuple(warnings))


def _length(pieces):
    return sum(len(p.text) for p in pieces) + max(0, len(pieces) - 1) * 2


def _tail(pieces, budget):
    result = []
    for piece in reversed(pieces):
        available = budget - _length(result) - (2 if result else 0)
        if available <= 0:
            break
        result.insert(0, replace(piece, text=piece.text[-available:]))
        if len(piece.text) > available:
            break
    return result


def _chunk(pieces, filename):
    return SourceChunk(
        text="\n\n".join(p.text for p in pieces), filename=filename,
        pages=tuple(dict.fromkeys(page for p in pieces for page in p.pages)),
        locations=tuple(dict.fromkeys(location for p in pieces for location in p.locations)),
    )


def chunk_document(document: Document, chunk_size=400, overlap=60):
    if chunk_size < 4 or not 0 <= overlap < chunk_size - 2:
        raise ValueError("分块长度至少为 4，重叠长度必须小于分块长度减 2。")
    pieces = []
    for passage in document.passages:
        if not passage.text.strip():
            continue
        if len(passage.text) <= chunk_size:
            pieces.append(passage)
            continue
        # Only oversized paragraphs are split further, preferring sentence ends.
        for sentence in re.split(r"(?<=[。！？；.!?;])\s*|\n+", passage.text):
            if not sentence:
                continue
            step = chunk_size - overlap - 2 if len(sentence) > chunk_size else chunk_size
            pieces.extend(replace(passage, text=sentence[i:i + step])
                          for i in range(0, len(sentence), step))
    chunks, current = [], []
    for piece in pieces:
        if current and _length(current + [piece]) > chunk_size:
            chunks.append(_chunk(current, document.filename))
            current = _tail(current, min(overlap, chunk_size - len(piece.text) - 2))
        current.append(piece)
    if current:
        chunks.append(_chunk(current, document.filename))
    return chunks


def unique_chunks(chunks):
    """Merge locations for identical text, so repeated pages remain traceable."""
    result = {}
    for chunk in chunks:
        old = result.get(chunk.text)
        result[chunk.text] = chunk if old is None else replace(
            old, pages=tuple(dict.fromkeys(old.pages + chunk.pages)),
            locations=tuple(dict.fromkeys(old.locations + chunk.locations)))
    return list(result.values())
