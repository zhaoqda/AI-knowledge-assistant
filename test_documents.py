from io import BytesIO
import unittest
import json
from unittest.mock import Mock, patch

from documents import Document, Passage, chunk_document, read_document, text_document, unique_chunks


def sample_pdf():
    """In-memory parser fixture: blank physical page 1, text on page 2."""
    from pypdf import PdfWriter
    from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
    writer = PdfWriter()
    writer.add_blank_page(width=300, height=300)
    page = writer.add_blank_page(width=300, height=300)
    font = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
                             NameObject("/Subtype"): NameObject("/Type1"),
                             NameObject("/BaseFont"): NameObject("/Helvetica")})
    page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})})
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 20 250 Td (Probation lasts three months.) Tj ET")
    page[NameObject("/Contents")] = stream
    buffer = BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


class DocumentTests(unittest.TestCase):
    def test_pdf_preserves_physical_page_after_blank(self):
        doc = read_document(sample_pdf(), "policy.PDF")
        self.assertEqual(doc.passages[0].pages, (2,))
        self.assertIn("three months", doc.text)
        self.assertIn("第 1 页", doc.warnings[0])

    def test_txt_bom_and_paragraphs(self):
        doc = read_document("\ufeff第一段\r\n\r\n第二段".encode(), "资料.txt")
        self.assertEqual([p.text for p in doc.passages], ["第一段", "第二段"])
        self.assertEqual(doc.passages[1].locations, ("第 2 段",))
        self.assertEqual(doc.passages[1].pages, ())

    def test_docx_preserves_body_table_order(self):
        import docx
        doc = docx.Document()
        doc.add_paragraph("第一段")
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "年假"
        table.cell(0, 1).text = "十天"
        doc.add_paragraph("末段")
        buffer = BytesIO()
        doc.save(buffer)
        result = read_document(buffer.getvalue(), "policy.docx")
        self.assertEqual([p.text for p in result.passages], ["第一段", "年假 | 十天", "末段"])
        self.assertEqual(result.passages[1].locations, ("表格 1 第 1 行",))

    def test_corrupt_file_has_readable_error(self):
        with self.assertRaises(ValueError):
            read_document(b"not a word document", "bad.docx")

    def test_txt_encoding_error(self):
        with self.assertRaisesRegex(ValueError, "UTF-8"):
            read_document(b"\xff", "bad.txt")

    def test_empty_text_returns_no_chunks(self):
        self.assertEqual(chunk_document(read_document(b"  ", "empty.txt")), [])

    def test_encrypted_pdf_requests_unlock(self):
        from pypdf import PdfWriter
        writer = PdfWriter()
        writer.add_blank_page(width=100, height=100)
        writer.encrypt("secret")
        buffer = BytesIO()
        writer.write(buffer)
        with self.assertRaisesRegex(ValueError, "解锁"):
            read_document(buffer.getvalue(), "locked.pdf")


class ChunkTests(unittest.TestCase):
    def test_rule_and_exception_paragraph_kept_whole(self):
        policy = "正式员工每周可以远程办公两天。试用期员工不适用此政策。"
        chunks = chunk_document(text_document("介绍" * 185 + "\n\n" + policy))
        self.assertTrue(any(policy in c.text for c in chunks))
        self.assertTrue(all(len(c.text) <= 400 for c in chunks))

    def test_overlap_carries_original_page_metadata(self):
        doc = Document("policy.pdf", (Passage("甲" * 350, pages=(1,)), Passage("乙" * 100, pages=(2,))))
        chunks = chunk_document(doc)
        self.assertEqual(len(chunks), 2)
        self.assertTrue(chunks[1].text.startswith("甲" * 60))
        self.assertEqual(chunks[1].pages, (1, 2))

    def test_missing_page_not_invented(self):
        doc = Document("policy.pdf", (Passage("第一部分", pages=(1,)), Passage("第三部分", pages=(3,))))
        self.assertEqual(chunk_document(doc)[0].pages, (1, 3))

    def test_long_sentence_has_bounded_chunks_and_overlap(self):
        text = "ABCDEFGHIJKLMNOPQRSTUVWXYZ" * 50
        chunks = chunk_document(text_document(text))
        self.assertTrue(all(len(c.text) <= 400 for c in chunks))
        rebuilt = chunks[0].text
        for chunk in chunks[1:]:
            # The shared prefix is followed by the next uninterrupted piece.
            prefix, remainder = chunk.text.split("\n\n", 1)
            self.assertTrue(rebuilt.endswith(prefix))
            rebuilt += remainder
        self.assertEqual(rebuilt, text)

    def test_sentence_boundaries_used_for_long_paragraph(self):
        sentence = "规则说明" * 20 + "。"
        chunks = chunk_document(text_document(sentence * 10), overlap=0)
        self.assertTrue(all(c.text.endswith("。") for c in chunks))

    def test_duplicate_text_merges_pages(self):
        doc = Document("policy.pdf", (Passage("甲" * 400, pages=(1,)), Passage("甲" * 400, pages=(2,))))
        chunks = unique_chunks(chunk_document(doc))
        self.assertEqual(len(chunks), 1)
        self.assertEqual(chunks[0].pages, (1, 2))

    def test_invalid_parameters(self):
        for size, overlap in [(0, 0), (100, 100), (100, -1)]:
            with self.assertRaises(ValueError):
                chunk_document(text_document("文字"), size, overlap)


class CitationTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        with patch("dotenv.load_dotenv"):
            import rag_core
        cls.core = rag_core

    def test_prompt_and_result_share_source_location(self):
        doc = Document("手册.pdf", (Passage("试用期为三个月。", pages=(4,)),))
        with patch.object(self.core, "call_llm", return_value=json.dumps({"answerable": True, "claims": [{"text": "三个月。", "source_id": 1, "quote": "试用期为三个月。"}]})) as model:
            answer, sources = self.core.query_document(doc, "试用期多久？")
        self.assertEqual(sources[0].pages, (4,))
        self.assertIn("[来源1] 手册.pdf · PDF 第 4 页", model.call_args.args[0])
        self.assertIn("[来源1]", answer)

    def test_cached_text_uses_current_filename_and_pages(self):
        retriever = Mock()
        retriever.search.return_value = [("试用期为三个月。", 0.8)]
        with patch.object(self.core, "call_llm", return_value=json.dumps({"answerable": True, "claims": [{"text": "三个月", "source_id": 1, "quote": "试用期为三个月。"}]})):
            for name, page in [("旧.pdf", 1), ("新.pdf", 9)]:
                doc = Document(name, (Passage("试用期为三个月。", pages=(page,)),))
                _, sources = self.core.query_document(doc, "多久转正？", retriever)
                self.assertEqual((sources[0].filename, sources[0].pages), (name, (page,)))


if __name__ == "__main__":
    unittest.main()
